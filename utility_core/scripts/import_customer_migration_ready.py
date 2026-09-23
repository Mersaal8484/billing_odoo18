#!/usr/bin/env python3
"""Safely stage and process the prepared customer migration workbook.

The script is read-only unless ``--apply`` is supplied to ``stage`` or
``process``.  It never creates transformers or feeders.  Customer rows are
accepted only when all mapping codes resolve and an already existing
transformer can be identified deterministically.

Typical use on this project::

    python utility_core/scripts/import_customer_migration_ready.py audit \
      --config invoice_utility_erp.conf --db invoice_utility_erp \
      --workbook "C:/Users/TUF/Downloads/البيانات/customer_migration_ready.xlsx" \
      --mapping "C:/Users/TUF/Downloads/جدول ترميز البيانات (utility.migration.mapping).xlsx" \
      --source-csv "C:/Users/TUF/Downloads/x_bi_sql_view.final_inv_pay2_(1).csv" \
      --default-category c1 --default-subscriber s1 \
      --default-contract TM1 --reading-date 2026-09-01

Run the same command with ``stage --apply`` after reviewing the audit, then
run ``process --apply --max-runs 1`` repeatedly.  Use ``reconcile`` for the
final count, network, reading, and balance checks.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Optional, Sequence, Tuple


DEFAULT_SHEET = "بيانات التهيئة"
HEADER_ROW = 4
DATA_ROW = 5

MAPPING_TYPES = {
    "المنطقة": "region",
    "الفرع": "area",
    "الفئة": "category",
    "نوع المشترك": "subscriber",
    "قالب العقد": "contract",
}

READY_COLUMNS = {
    "name": 0,
    "mobile": 1,
    "national_id": 2,
    "customer_number": 3,
    "subscriber_no": 4,
    "char_code": 5,
    "is_active": 6,
    "region": 7,
    "area": 8,
    "category": 9,
    "subscriber": 10,
    "contract": 11,
    "meter_number": 12,
    "meter_reading": 13,
    "opening_reading": 14,
    "previous_balance": 15,
    "current_balance": 16,
    "phase": 17,
    "is_private_transformer": 18,
    "owner_reference": 19,
    "transformer_alias": 20,
    "current_reading": 21,
    "last_reading_date": 22,
}

# Columns in x_bi_sql_view.final_inv_pay2_(1).csv.
SOURCE_CUSTOMER_NUMBER = 0
SOURCE_NAME = 1
SOURCE_METER_NUMBER = 3
SOURCE_MOBILE = 17
SOURCE_TRANSFORMER_METER = 36
SOURCE_TRANSFORMER_NAME = 51
SOURCE_TRANSFORMER_ID = 52
SOURCE_CURRENT_READING = 53
SOURCE_LAST_INVOICE_READING = 54


def clean(value) -> str:
    return "" if value is None else str(value).strip()


def numeric(value) -> float:
    text = clean(value).replace(",", "").replace("٬", "").replace("٫", ".")
    if text.startswith("'"):
        text = text[1:]
    return float(text or 0.0)


def meter_match_key(value) -> str:
    """Normalize only for matching; the exact source meter remains canonical."""
    value = clean(value)
    if value.isdigit():
        return value.lstrip("0") or "0"
    return value


def normalize_text(value) -> str:
    value = clean(value).lower()
    value = value.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    value = value.replace("ة", "ه").replace("ى", "ي")
    return re.sub(r"[^\w\u0600-\u06ff]+", " ", value).strip()


def parse_bool(value, default: bool) -> bool:
    value = clean(value).lower()
    if not value:
        return default
    if value in {"نعم", "yes", "true", "1", "y", "t", "صح"}:
        return True
    if value in {"لا", "no", "false", "0", "n", "f", "خطأ"}:
        return False
    raise ValueError("INVALID_BOOLEAN_VALUE: %r" % value)


def parse_phase(value) -> str:
    value = clean(value).lower()
    if value in {"single", "1", "1 phase", "single phase", "واحد فاز", "فاز واحد", "أحادي", "احادي"}:
        return "single"
    if value in {"three", "3", "3 phase", "three phase", "ثلاثة فاز", "ثلاثي", "محول تيار"}:
        return "three"
    raise ValueError("INVALID_METER_PHASE: %r" % value)


def parse_date(value, default_value: Optional[date]) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean(value)
    if not text:
        return default_value
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError("INVALID_DATE_VALUE: %r" % value)


def bracket_tokens(value) -> Iterable[str]:
    for token in re.findall(r"\[([^\]]+)\]", clean(value)):
        if token.strip():
            yield token.strip()


@dataclass
class InputRow:
    workbook_row: int
    values: List[object]
    source_row: Optional[int]
    source: Optional[List[str]]

    def get(self, field: str):
        index = READY_COLUMNS[field]
        return self.values[index] if index < len(self.values) else None


@dataclass
class PreparedRow:
    workbook_row: int
    values: Dict[str, object]
    errors: List[str]


@dataclass
class Preparation:
    queued: List[PreparedRow]
    errors: List[PreparedRow]
    skipped: Counter
    warnings: Counter
    transformer_aliases: Dict[str, str]
    unresolved_aliases: Counter


class MappingRegistry:
    def __init__(self):
        self.aliases: Dict[str, Dict[str, str]] = defaultdict(dict)

    def add(self, mapping_type: str, legacy_code: str, target_value) -> None:
        legacy_code = clean(legacy_code)
        if not legacy_code:
            return
        candidates = {legacy_code, clean(target_value), normalize_text(target_value)}
        candidates.update(bracket_tokens(target_value))
        for candidate in candidates:
            candidate = clean(candidate)
            if not candidate:
                continue
            self.aliases[mapping_type][candidate] = legacy_code
            self.aliases[mapping_type][candidate.lower()] = legacy_code

    def resolve(self, mapping_type: str, value, default: Optional[str]) -> Optional[str]:
        text = clean(value)
        if not text:
            text = clean(default)
        if not text:
            return None
        aliases = self.aliases[mapping_type]
        for candidate in (text, text.lower(), normalize_text(text)):
            if candidate in aliases:
                return aliases[candidate]
        return None


def load_mapping_registry(path: Path) -> MappingRegistry:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    registry = MappingRegistry()
    sheet = workbook.active
    for row in sheet.iter_rows(min_row=2, values_only=True):
        mapping_type = MAPPING_TYPES.get(clean(row[1]))
        if not mapping_type:
            continue
        legacy_code = clean(row[2])
        target_index = {
            "region": 3,
            "area": 4,
            "category": 5,
            "subscriber": 6,
            "contract": 7,
        }[mapping_type]
        registry.add(mapping_type, legacy_code, row[target_index])
    workbook.close()
    return registry


def load_source_rows(path: Path) -> Dict[Tuple[str, str], List[Tuple[int, List[str]]]]:
    by_pair: Dict[Tuple[str, str], List[Tuple[int, List[str]]]] = defaultdict(list)
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row_number, row in enumerate(csv.reader(handle), start=1):
            if row_number == 1:
                continue
            key = (clean(row[SOURCE_CUSTOMER_NUMBER]), meter_match_key(row[SOURCE_METER_NUMBER]))
            by_pair[key].append((row_number, row))
    return by_pair


def choose_source(
    candidates: Sequence[Tuple[int, List[str]]], exact_meter: str
) -> Tuple[Optional[int], Optional[List[str]]]:
    exact = [item for item in candidates if clean(item[1][SOURCE_METER_NUMBER]) == exact_meter]
    if len(exact) == 1:
        return exact[0]
    if len(candidates) == 1:
        return candidates[0]
    return None, None


def load_input_rows(workbook_path: Path, sheet_name: str, source_csv: Path) -> List[InputRow]:
    import openpyxl

    source_index = load_source_rows(source_csv)
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    if sheet_name not in workbook.sheetnames:
        raise RuntimeError("Worksheet not found: %s" % sheet_name)
    sheet = workbook[sheet_name]
    rows: List[InputRow] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=DATA_ROW, values_only=True), start=DATA_ROW):
        values = list(row)
        if not any(clean(value) for value in values):
            continue
        customer_number = clean(values[READY_COLUMNS["customer_number"]])
        ready_meter = clean(values[READY_COLUMNS["meter_number"]])
        source_row, source = choose_source(
            source_index.get((customer_number, meter_match_key(ready_meter)), []),
            ready_meter,
        )
        rows.append(InputRow(row_number, values, source_row, source))
    workbook.close()
    return rows


def load_explicit_transformer_map(path: Optional[Path]) -> Dict[str, str]:
    if not path:
        return {}
    result = {}
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        required = {"prepared_code", "odoo_code"}
        if not required.issubset(reader.fieldnames or []):
            raise RuntimeError("Transformer map requires prepared_code,odoo_code columns")
        for row in reader:
            prepared_code = clean(row["prepared_code"])
            odoo_code = clean(row["odoo_code"])
            if prepared_code and odoo_code:
                result[prepared_code] = odoo_code
    return result


def _unique_alias(alias_map: Dict[str, set], alias: str, target: str) -> None:
    alias = clean(alias)
    target = clean(target)
    if alias and target:
        alias_map.setdefault(alias, set()).add(target)


def resolve_transformer_aliases(env, rows: Sequence[InputRow], company_id: int, explicit: Mapping[str, str]):
    transformers = env["utility.transformer"].search([("company_id", "=", company_id)])
    by_code = {record.code: record for record in transformers}
    alias_candidates: Dict[str, set] = {}

    for code in by_code:
        _unique_alias(alias_candidates, code, code)
    for prepared_code, odoo_code in explicit.items():
        if odoo_code in by_code:
            _unique_alias(alias_candidates, prepared_code, odoo_code)

    transformer_staging = env["utility.migration.transformer"].search([
        ("company_id", "=", company_id),
        ("created_transformer_id", "!=", False),
    ])
    for staging in transformer_staging:
        target_code = staging.created_transformer_id.code
        for alias in (staging.reference, staging.transformer_code, staging.legacy_analytic_id):
            _unique_alias(alias_candidates, alias, target_code)

    normalized_names = {
        record.code: normalize_text(record.name)
        for record in transformers
        if normalize_text(record.name)
    }
    alias_sources: Dict[str, List[InputRow]] = defaultdict(list)
    for row in rows:
        alias_sources[clean(row.get("transformer_alias"))].append(row)

    for alias, source_rows in alias_sources.items():
        if not alias or len(alias_candidates.get(alias, set())) == 1:
            continue
        candidates = set()
        evidence_values = set()
        for item in source_rows:
            if not item.source:
                continue
            source_name = normalize_text(item.source[SOURCE_TRANSFORMER_NAME])
            source_meter = clean(item.source[SOURCE_TRANSFORMER_METER])
            if len(source_name) >= 6:
                evidence_values.add(source_name)
            if len(source_meter) >= 6:
                evidence_values.add(source_meter)
        for code, transformer_name in normalized_names.items():
            if any(
                evidence == transformer_name
                or evidence in transformer_name
                or transformer_name in evidence
                for evidence in evidence_values
                if len(evidence) >= 6
            ):
                candidates.add(code)
        if len(candidates) == 1:
            alias_candidates[alias] = candidates

    resolved = {
        alias: next(iter(candidates))
        for alias, candidates in alias_candidates.items()
        if len(candidates) == 1
    }
    ambiguous = {
        alias: sorted(candidates)
        for alias, candidates in alias_candidates.items()
        if len(candidates) > 1
    }
    input_aliases = {clean(row.get("transformer_alias")) for row in rows}
    resolved = {alias: code for alias, code in resolved.items() if alias in input_aliases}
    ambiguous = {alias: codes for alias, codes in ambiguous.items() if alias in input_aliases}
    return resolved, ambiguous, by_code


def validate_database_mappings(env, company_id: int, registry: MappingRegistry) -> set:
    cache = env["utility.migration.mapping"].get_mapping_cache(company_id)
    return set(cache)


def prepare(env, args, input_rows: Sequence[InputRow], registry: MappingRegistry) -> Preparation:
    company_id = args.company_id
    explicit = load_explicit_transformer_map(args.transformer_map)
    transformer_aliases, ambiguous_aliases, transformers = resolve_transformer_aliases(
        env, input_rows, company_id, explicit
    )
    db_mapping_keys = validate_database_mappings(env, company_id, registry)

    env.cr.execute(
        "SELECT customer_number FROM utility_customer WHERE company_id=%s",
        (company_id,),
    )
    existing_customers = {clean(row[0]) for row in env.cr.fetchall()}
    env.cr.execute(
        "SELECT meter_number FROM utility_meter WHERE company_id=%s",
        (company_id,),
    )
    existing_meters = {clean(row[0]) for row in env.cr.fetchall()}
    env.cr.execute(
        "SELECT customer_number FROM utility_migration_customer WHERE company_id=%s",
        (company_id,),
    )
    existing_staging = {clean(row[0]) for row in env.cr.fetchall()}

    customer_counts = Counter(clean(row.get("customer_number")) for row in input_rows)
    canonical_meters = [
        clean(row.source[SOURCE_METER_NUMBER]) if row.source else clean(row.get("meter_number"))
        for row in input_rows
    ]
    meter_counts = Counter(canonical_meters)

    queued: List[PreparedRow] = []
    errors: List[PreparedRow] = []
    skipped = Counter()
    warnings = Counter()
    unresolved_aliases = Counter()

    defaults = {
        "region": args.default_region,
        "area": args.default_area,
        "category": args.default_category,
        "subscriber": args.default_subscriber,
        "contract": args.default_contract,
    }
    reading_date_default = parse_date(args.reading_date, None)

    for row, canonical_meter in zip(input_rows, canonical_meters):
        customer_number = clean(row.get("customer_number"))
        if customer_number in existing_customers:
            skipped["CUSTOMER_ALREADY_EXISTS"] += 1
            continue
        if canonical_meter in existing_meters:
            skipped["METER_ALREADY_EXISTS"] += 1
            continue
        if customer_number in existing_staging:
            skipped["STAGING_ALREADY_EXISTS"] += 1
            continue

        row_errors = []
        if not customer_number:
            row_errors.append("MISSING_CUSTOMER_NUMBER")
        if not canonical_meter:
            row_errors.append("MISSING_METER_NUMBER")
        if customer_number and customer_counts[customer_number] > 1:
            row_errors.append("DUPLICATE_CUSTOMER_NUMBER_IN_WORKBOOK")
        if canonical_meter and meter_counts[canonical_meter] > 1:
            row_errors.append("DUPLICATE_METER_NUMBER_IN_WORKBOOK")
        if not row.source:
            row_errors.append("SOURCE_ROW_NOT_FOUND_OR_AMBIGUOUS")

        legacy_codes = {}
        for mapping_type, workbook_field in (
            ("region", "region"),
            ("area", "area"),
            ("category", "category"),
            ("subscriber", "subscriber"),
            ("contract", "contract"),
        ):
            legacy_code = registry.resolve(
                mapping_type, row.get(workbook_field), defaults[mapping_type]
            )
            legacy_codes[mapping_type] = legacy_code
            if not legacy_code:
                row_errors.append("UNRESOLVED_%s_MAPPING" % mapping_type.upper())
            elif (mapping_type, legacy_code) not in db_mapping_keys:
                row_errors.append("MAPPING_NOT_LOADED_IN_ODOO:%s/%s" % (mapping_type, legacy_code))

        transformer_alias = clean(row.get("transformer_alias"))
        transformer_code = transformer_aliases.get(transformer_alias)
        if transformer_alias in ambiguous_aliases:
            row_errors.append("AMBIGUOUS_TRANSFORMER_ALIAS:%s" % transformer_alias)
        elif not transformer_code or transformer_code not in transformers:
            row_errors.append("MISSING_EXISTING_TRANSFORMER:%s" % (transformer_alias or "<blank>"))
            unresolved_aliases[transformer_alias or "<blank>"] += 1

        try:
            phase = parse_phase(row.get("phase"))
        except ValueError as error:
            row_errors.append(str(error))
            phase = False
        try:
            active = parse_bool(row.get("is_active"), True)
            private = parse_bool(row.get("is_private_transformer"), False)
        except ValueError as error:
            row_errors.append(str(error))
            active, private = True, False
        try:
            last_reading_date = parse_date(row.get("last_reading_date"), reading_date_default)
        except ValueError as error:
            row_errors.append(str(error))
            last_reading_date = None
        if not last_reading_date:
            row_errors.append("MISSING_LAST_READING_DATE")

        source = row.source
        opening_reading = (
            numeric(source[SOURCE_LAST_INVOICE_READING])
            if source else numeric(row.get("opening_reading"))
        )
        current_reading = (
            numeric(source[SOURCE_CURRENT_READING])
            if source else numeric(row.get("current_reading"))
        )
        if current_reading < opening_reading:
            warnings["CURRENT_READING_BELOW_LAST_INVOICE_READING"] += 1
            if not args.allow_reading_regression:
                row_errors.append("CURRENT_READING_BELOW_LAST_INVOICE_READING")

        values = {
            "company_id": company_id,
            "name": clean(row.get("name")) or (clean(source[SOURCE_NAME]) if source else ""),
            "mobile": clean(row.get("mobile")) or (clean(source[SOURCE_MOBILE]) if source else ""),
            "national_id": clean(row.get("national_id")) or False,
            "customer_number": customer_number or "__MISSING_ROW_%s" % row.workbook_row,
            "subscriber_no": clean(row.get("subscriber_no")) or customer_number,
            "char_code": clean(row.get("char_code")) or False,
            "is_active": active,
            "legacy_region": legacy_codes["region"] or False,
            "legacy_area": legacy_codes["area"] or False,
            "legacy_category": legacy_codes["category"] or False,
            "legacy_subscriber_type": legacy_codes["subscriber"] or False,
            "legacy_contract": legacy_codes["contract"] or False,
            "legacy_transformer_code": transformer_code or False,
            "meter_number": canonical_meter or False,
            "meter_reading": int(opening_reading),
            "last_reading": opening_reading,
            "opening_reading": opening_reading,
            "current_reading": current_reading,
            "last_reading_date": last_reading_date,
            "previous_balance": "%.2f" % numeric(row.get("previous_balance")),
            "current_balance": numeric(row.get("current_balance")),
            "phase": phase,
            "is_private_transformer": private,
            "owner_reference": clean(row.get("owner_reference")) or False,
            "source_row_number": row.workbook_row,
        }
        prepared = PreparedRow(row.workbook_row, values, row_errors)
        if row_errors:
            values.update({"state": "error", "error_message": "; ".join(row_errors)})
            errors.append(prepared)
        else:
            values["state"] = "queued"
            queued.append(prepared)

    return Preparation(
        queued=queued,
        errors=errors,
        skipped=skipped,
        warnings=warnings,
        transformer_aliases=transformer_aliases,
        unresolved_aliases=unresolved_aliases,
    )


def load_odoo(args):
    odoo_root = str(args.odoo_root.resolve())
    if odoo_root not in sys.path:
        sys.path.insert(0, odoo_root)
    import odoo
    from odoo import api
    from odoo.modules.registry import Registry

    odoo.tools.config.parse_config([
        "-c", str(args.config.resolve()),
        "-d", args.db,
        "--no-http",
        "--log-level=warn",
    ])
    return odoo, api, Registry(args.db)


def job_reference(args) -> str:
    if args.job_reference:
        return args.job_reference
    digest = hashlib.sha256()
    for path in (args.workbook, args.mapping, args.source_csv, args.transformer_map):
        if path:
            digest.update(path.read_bytes())
    digest.update(json.dumps({
        "company": args.company_id,
        "region": args.default_region,
        "area": args.default_area,
        "category": args.default_category,
        "subscriber": args.default_subscriber,
        "contract": args.default_contract,
        "reading_date": args.reading_date,
    }, sort_keys=True).encode("utf-8"))
    return "customer-ready-%s" % digest.hexdigest()[:16]


def report_preparation(preparation: Preparation, reference: str) -> None:
    print("JOB_REFERENCE=%s" % reference)
    print("QUEUED=%s" % len(preparation.queued))
    print("ERRORS=%s" % len(preparation.errors))
    print("SKIPPED=%s" % sum(preparation.skipped.values()))
    print("SKIPPED_REASONS=%s" % dict(preparation.skipped))
    print("WARNINGS=%s" % dict(preparation.warnings))
    print("RESOLVED_TRANSFORMER_ALIASES=%s" % len(preparation.transformer_aliases))
    print("UNRESOLVED_TRANSFORMER_ALIASES=%s" % len(preparation.unresolved_aliases))
    if preparation.unresolved_aliases:
        print("UNRESOLVED_ALIAS_ROWS=%s" % preparation.unresolved_aliases.most_common(30))
    error_reasons = Counter()
    for row in preparation.errors:
        error_reasons.update(row.errors)
    print("ERROR_REASONS=%s" % error_reasons.most_common())
    print("QUEUED_PREVIOUS_BALANCE=%.2f" % sum(
        numeric(row.values["previous_balance"]) for row in preparation.queued
    ))
    print("QUEUED_CURRENT_BALANCE=%.2f" % sum(
        numeric(row.values["current_balance"]) for row in preparation.queued
    ))


def export_unresolved_transformer_map(
    path: Optional[Path], preparation: Preparation, rows: Sequence[InputRow]
) -> None:
    if not path:
        return
    evidence = defaultdict(lambda: {
        "ids": set(), "names": set(), "meters": set(), "rows": 0,
    })
    for row in rows:
        alias = clean(row.get("transformer_alias")) or "<blank>"
        if alias not in preparation.unresolved_aliases:
            continue
        evidence[alias]["rows"] += 1
        if row.source:
            evidence[alias]["ids"].add(clean(row.source[SOURCE_TRANSFORMER_ID]))
            evidence[alias]["names"].add(clean(row.source[SOURCE_TRANSFORMER_NAME]))
            evidence[alias]["meters"].add(clean(row.source[SOURCE_TRANSFORMER_METER]))
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=[
            "prepared_code", "row_count", "source_transformer_ids",
            "source_transformer_names", "source_transformer_meters", "odoo_code",
        ])
        writer.writeheader()
        for alias in sorted(evidence):
            item = evidence[alias]
            writer.writerow({
                "prepared_code": alias,
                "row_count": item["rows"],
                "source_transformer_ids": " | ".join(sorted(item["ids"])),
                "source_transformer_names": " | ".join(sorted(item["names"])),
                "source_transformer_meters": " | ".join(sorted(item["meters"])),
                "odoo_code": "",
            })
    print("UNRESOLVED_TRANSFORMER_MAP_WRITTEN=%s" % path.resolve())


def command_audit(args) -> None:
    rows = load_input_rows(args.workbook, args.sheet, args.source_csv)
    registry = load_mapping_registry(args.mapping)
    odoo, api, odoo_registry = load_odoo(args)
    with odoo_registry.cursor() as cursor:
        env = api.Environment(cursor, odoo.SUPERUSER_ID, {
            "allowed_company_ids": [args.company_id],
            "company_id": args.company_id,
        })
        preparation = prepare(env, args, rows, registry)
        cursor.rollback()
    print("WORKBOOK_ROWS=%s" % len(rows))
    report_preparation(preparation, job_reference(args))
    export_unresolved_transformer_map(args.unresolved_map_output, preparation, rows)


def command_stage(args) -> None:
    rows = load_input_rows(args.workbook, args.sheet, args.source_csv)
    registry = load_mapping_registry(args.mapping)
    reference = job_reference(args)
    odoo, api, odoo_registry = load_odoo(args)
    with odoo_registry.cursor() as cursor:
        env = api.Environment(cursor, odoo.SUPERUSER_ID, {
            "allowed_company_ids": [args.company_id],
            "company_id": args.company_id,
        })
        existing_batch = env["utility.migration.batch"].search([
            ("company_id", "=", args.company_id),
            ("job_reference", "=", reference),
        ], limit=1)
        if existing_batch:
            print("BATCH_ALREADY_EXISTS=%s STATE=%s" % (existing_batch.name, existing_batch.state))
            cursor.rollback()
            return
        preparation = prepare(env, args, rows, registry)
        report_preparation(preparation, reference)
        export_unresolved_transformer_map(args.unresolved_map_output, preparation, rows)
        if not args.apply:
            cursor.rollback()
            print("DRY_RUN=1 (add --apply to create the staging batch)")
            return
        if preparation.errors and not args.allow_errors:
            cursor.rollback()
            raise RuntimeError(
                "%s invalid rows remain. Resolve them first or explicitly add --allow-errors."
                % len(preparation.errors)
            )
        if not preparation.queued:
            cursor.rollback()
            raise RuntimeError("No valid rows are available for staging")
        batch = env["utility.migration.batch"].create({
            "company_id": args.company_id,
            "migration_type": "customer",
            "state": "queued",
            "chunk_size": args.chunk_size,
            "job_reference": reference,
        })
        all_rows = preparation.queued + preparation.errors
        values = []
        for row in all_rows:
            payload = dict(row.values)
            payload["last_batch_id"] = batch.id
            values.append(payload)
        staging_model = env["utility.migration.customer"]
        for start in range(0, len(values), 500):
            staging_model.create(values[start:start + 500])
        cursor.commit()
        print("BATCH_CREATED=%s ID=%s TOTAL=%s" % (batch.name, batch.id, len(values)))


def find_batch(env, args):
    reference = job_reference(args)
    batch = env["utility.migration.batch"].search([
        ("company_id", "=", args.company_id),
        ("job_reference", "=", reference),
    ], limit=1)
    if not batch:
        raise RuntimeError("Migration batch not found for %s" % reference)
    return batch


def fix_opening_reading_dates(cursor, batch_id: int, superuser_id: int) -> int:
    cursor.execute(
        """
        UPDATE utility_reading reading
           SET reading_date=staging.last_reading_date,
               write_date=now(),
               write_uid=%s
          FROM utility_migration_customer staging
         WHERE staging.last_batch_id=%s
           AND staging.state='imported'
           AND staging.last_reading_date IS NOT NULL
           AND staging.created_reading_id=reading.id
           AND reading.reading_source='legacy_migration'
           AND reading.reading_purpose='opening'
           AND reading.reading_date::date<>staging.last_reading_date
        """,
        (superuser_id, batch_id),
    )
    return cursor.rowcount


def command_process(args) -> None:
    odoo, api, odoo_registry = load_odoo(args)
    with odoo_registry.cursor() as cursor:
        env = api.Environment(cursor, odoo.SUPERUSER_ID, {
            "allowed_company_ids": [args.company_id],
            "company_id": args.company_id,
        })
        batch = find_batch(env, args)
        if not args.apply:
            print("BATCH=%s STATE=%s TOTAL=%s PROCESSED=%s SUCCESS=%s ERROR=%s" % (
                batch.name, batch.state, batch.record_count, batch.processed_count,
                batch.success_count, batch.error_count,
            ))
            cursor.rollback()
            print("DRY_RUN=1 (add --apply to process queued rows)")
            return
        run_number = 0
        while batch.state not in ("done", "partial", "cancelled", "failed"):
            if args.max_runs and run_number >= args.max_runs:
                break
            batch.action_process_batch(max_records_per_run=args.run_size)
            fixed_dates = fix_opening_reading_dates(cursor, batch.id, odoo.SUPERUSER_ID)
            cursor.commit()
            env.clear()
            batch = env["utility.migration.batch"].browse(batch.id)
            run_number += 1
            print("RUN=%s BATCH=%s STATE=%s PROCESSED=%s SUCCESS=%s ERROR=%s FIXED_DATES=%s" % (
                run_number, batch.name, batch.state, batch.processed_count,
                batch.success_count, batch.error_count, fixed_dates,
            ))


def command_reconcile(args) -> None:
    odoo, api, odoo_registry = load_odoo(args)
    with odoo_registry.cursor() as cursor:
        env = api.Environment(cursor, odoo.SUPERUSER_ID, {
            "allowed_company_ids": [args.company_id],
            "company_id": args.company_id,
        })
        batch = find_batch(env, args)
        cursor.execute(
            """
            WITH receivable AS (
                SELECT line.move_id, sum(line.balance) AS balance
                  FROM account_move_line line
                  JOIN account_account account ON account.id=line.account_id
                 WHERE account.account_type='asset_receivable'
              GROUP BY line.move_id
            )
            SELECT
                count(*) AS total,
                count(*) FILTER (WHERE staging.state='imported') AS imported,
                count(*) FILTER (WHERE staging.state='error') AS errors,
                count(*) FILTER (WHERE staging.state IN ('queued','processing')) AS pending,
                count(*) FILTER (
                    WHERE staging.state='imported'
                      AND (customer.transformer_id IS NULL OR customer.route_id IS NULL OR route.transformer_id<>customer.transformer_id)
                ) AS bad_network,
                count(*) FILTER (
                    WHERE staging.state='imported' AND abs(meter.last_reading_value-staging.current_reading)>0.001
                ) AS bad_current_reading,
                count(*) FILTER (
                    WHERE staging.state='imported' AND meter.last_read_date::date<>staging.last_reading_date
                ) AS bad_reading_date,
                count(*) FILTER (
                    WHERE staging.state='imported' AND opening.reading_date::date<>staging.last_reading_date
                ) AS bad_opening_reading_date,
                count(*) FILTER (
                    WHERE staging.state='imported' AND abs(partner.open_balance-staging.current_balance)>0.01
                ) AS bad_current_balance,
                round(sum(staging.current_balance) FILTER (WHERE staging.state='imported')::numeric, 2) AS source_current_balance,
                round(sum(partner.open_balance) FILTER (WHERE staging.state='imported')::numeric, 2) AS target_current_balance,
                round(sum(staging.previous_balance::numeric) FILTER (WHERE staging.state='imported'), 2) AS source_previous_balance,
                round(sum(receivable.balance) FILTER (WHERE staging.state='imported')::numeric, 2) AS target_previous_balance
            FROM utility_migration_customer staging
            LEFT JOIN utility_customer customer ON customer.id=staging.created_customer_id
            LEFT JOIN utility_route route ON route.id=customer.route_id
            LEFT JOIN utility_meter meter ON meter.id=staging.created_meter_id
            LEFT JOIN utility_reading opening ON opening.id=staging.created_reading_id
            LEFT JOIN res_partner partner ON partner.id=staging.created_partner_id
            LEFT JOIN receivable ON receivable.move_id=staging.opening_move_id
            WHERE staging.last_batch_id=%s
            """,
            (batch.id,),
        )
        columns = [description[0] for description in cursor.description]
        print("RECONCILIATION=%s" % dict(zip(columns, cursor.fetchone())))
        cursor.execute(
            """
            SELECT split_part(error_message, ';', 1), count(*)
              FROM utility_migration_customer
             WHERE last_batch_id=%s AND state='error'
          GROUP BY split_part(error_message, ';', 1)
          ORDER BY count(*) DESC
            """,
            (batch.id,),
        )
        print("ERRORS=%s" % cursor.fetchall())
        cursor.rollback()


def add_common_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--config", type=Path, required=True, help="Odoo configuration file")
    parser.add_argument("--db", required=True, help="Odoo database name")
    parser.add_argument(
        "--odoo-root", type=Path,
        default=Path(r"D:\odoo-16.0\odoo-16.0"),
        help="Directory containing odoo-bin and the odoo package",
    )
    parser.add_argument("--company-id", type=int, default=1)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--mapping", type=Path, required=True)
    parser.add_argument("--source-csv", type=Path, required=True)
    parser.add_argument("--transformer-map", type=Path)
    parser.add_argument(
        "--unresolved-map-output", type=Path,
        help="Optional CSV template containing unresolved transformer aliases and evidence",
    )
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--default-region")
    parser.add_argument("--default-area")
    parser.add_argument("--default-category")
    parser.add_argument("--default-subscriber")
    parser.add_argument("--default-contract")
    parser.add_argument("--reading-date", help="YYYY-MM-DD fallback for blank dates")
    parser.add_argument("--job-reference")
    parser.add_argument("--chunk-size", type=int, default=200)
    parser.add_argument("--run-size", type=int, default=1000)
    parser.add_argument("--max-runs", type=int, default=1, help="0 means process until terminal state")
    parser.add_argument("--apply", action="store_true", help="Required for database writes")
    parser.add_argument(
        "--allow-errors", action="store_true",
        help="Allow stage --apply to create error rows alongside valid queued rows",
    )
    parser.add_argument(
        "--allow-reading-regression", action="store_true",
        help="Allow a current reading below the last invoiced reading (rollover/manual review)",
    )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    for command in ("audit", "stage", "process", "reconcile"):
        subparser = subparsers.add_parser(command)
        add_common_arguments(subparser)
    return parser


def validate_paths(args) -> None:
    required = (args.config, args.odoo_root, args.workbook, args.mapping, args.source_csv)
    missing = [str(path) for path in required if not path.exists()]
    if args.transformer_map and not args.transformer_map.exists():
        missing.append(str(args.transformer_map))
    if missing:
        raise RuntimeError("Missing paths: %s" % ", ".join(missing))
    if args.command in {"audit", "stage"} and args.apply and args.command == "audit":
        raise RuntimeError("audit is always read-only; remove --apply")


def main(argv: Optional[Sequence[str]] = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = build_parser()
    args = parser.parse_args(argv)
    validate_paths(args)
    commands = {
        "audit": command_audit,
        "stage": command_stage,
        "process": command_process,
        "reconcile": command_reconcile,
    }
    commands[args.command](args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
